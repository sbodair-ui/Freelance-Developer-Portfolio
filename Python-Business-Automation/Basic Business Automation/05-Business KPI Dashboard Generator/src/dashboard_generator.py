from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from src.kpi_data_loader import load_business_data
from src.kpi_calculator import calculate_all_kpis
from src.trend_analyzer import analyze_all_trends
from src.chart_generator import create_all_charts


DATA_FILE = Path("sample_data/business_data.csv")
OUTPUT_DIRECTORY = Path("output")
DASHBOARD_FILE = OUTPUT_DIRECTORY / "dashboard.xlsx"
KPI_SUMMARY_FILE = OUTPUT_DIRECTORY / "kpi_summary.csv"


def create_kpi_summary(kpis):
    """
    Convert KPI results into a DataFrame.
    """
    return pd.DataFrame(
        [
            ["Total Revenue", kpis["total_revenue"]],
            ["Total Expenses", kpis["total_expenses"]],
            ["Net Revenue", kpis["net_revenue"]],
            ["Profit Margin", kpis["profit_margin"]],
            ["Total Units", kpis["total_units"]],
            ["Total Transactions", kpis["total_transactions"]],
            [
                "Average Transaction Value",
                kpis["average_transaction_value"],
            ],
        ],
        columns=["KPI", "Value"],
    )


def generate_excel_dashboard(kpis, analyses, output_file):
    """
    Generate the Excel business KPI dashboard.
    """
    output_file = Path(output_file)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    kpi_summary = create_kpi_summary(kpis)

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        kpi_summary.to_excel(
            writer,
            sheet_name="KPI Summary",
            index=False,
        )

        analyses["monthly_performance"].to_excel(
            writer,
            sheet_name="Monthly Performance",
            index=False,
        )

        analyses["department_performance"].to_excel(
            writer,
            sheet_name="Department Performance",
            index=False,
        )

        analyses["product_performance"].to_excel(
            writer,
            sheet_name="Product Performance",
            index=False,
        )

    format_excel_dashboard(output_file)

    return output_file


def format_excel_dashboard(output_file):
    """
    Apply basic professional formatting to the Excel dashboard.
    """
    workbook = load_workbook(output_file)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for worksheet in workbook.worksheets:
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        worksheet.freeze_panes = "A2"

        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value)),
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = min(max_length + 2, 30)

    # Format currency and percentage values.
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.column == 2 and isinstance(
                    cell.value,
                    (int, float),
                ):
                    if worksheet.title == "KPI Summary":
                        if cell.row == 5:
                            cell.number_format = "0.00%"
                        elif cell.row in [2, 3, 4, 8]:
                            cell.number_format = '$#,##0.00'

    workbook.save(output_file)


def save_kpi_summary(kpis, output_file):
    """
    Save KPI results to a CSV file.
    """
    summary = create_kpi_summary(kpis)
    summary.to_csv(output_file, index=False)

    return output_file


def generate_dashboard():
    """
    Run the complete business KPI dashboard workflow.
    """
    OUTPUT_DIRECTORY.mkdir(parents=True, exist_ok=True)

    print("Loading business data...")

    data = load_business_data(DATA_FILE)

    print(f"Loaded {len(data)} business records.")

    print("Calculating KPIs...")

    kpis = calculate_all_kpis(data)

    print("Analyzing business performance...")

    analyses = analyze_all_trends(data)

    print("Generating charts...")

    charts = create_all_charts(
        analyses["monthly_performance"],
        analyses["department_performance"],
        analyses["product_performance"],
        OUTPUT_DIRECTORY,
    )

    print("Generating Excel dashboard...")

    dashboard = generate_excel_dashboard(
        kpis,
        analyses,
        DASHBOARD_FILE,
    )

    print("Saving KPI summary...")

    kpi_summary = save_kpi_summary(
        kpis,
        KPI_SUMMARY_FILE,
    )

    print("\nDashboard generation complete!")
    print(f"Excel Dashboard: {dashboard}")
    print(f"KPI Summary: {kpi_summary}")

    print("\nGenerated charts:")

    for chart_name, chart_path in charts.items():
        print(f"- {chart_name}: {chart_path}")

    return {
        "dashboard": dashboard,
        "kpi_summary": kpi_summary,
        "charts": charts,
    }


if __name__ == "__main__":
    generate_dashboard()
