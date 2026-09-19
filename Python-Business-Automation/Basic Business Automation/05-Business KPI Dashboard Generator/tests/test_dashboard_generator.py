from pathlib import Path

from openpyxl import load_workbook

from src.kpi_data_loader import load_business_data
from src.kpi_calculator import calculate_all_kpis
from src.trend_analyzer import analyze_all_trends
from src.dashboard_generator import (
    create_kpi_summary,
    generate_excel_dashboard,
    save_kpi_summary,
)


DATA_FILE = "sample_data/business_data.csv"


def test_create_kpi_summary():
    data = load_business_data(DATA_FILE)
    kpis = calculate_all_kpis(data)

    summary = create_kpi_summary(kpis)

    assert len(summary) == 7
    assert "KPI" in summary.columns
    assert "Value" in summary.columns


def test_generate_excel_dashboard(tmp_path):
    data = load_business_data(DATA_FILE)
    kpis = calculate_all_kpis(data)
    analyses = analyze_all_trends(data)

    output_file = tmp_path / "dashboard.xlsx"

    result = generate_excel_dashboard(
        kpis,
        analyses,
        output_file,
    )

    assert Path(result).exists()

    workbook = load_workbook(result)

    assert workbook.sheetnames == [
        "KPI Summary",
        "Monthly Performance",
        "Department Performance",
        "Product Performance",
    ]


def test_save_kpi_summary(tmp_path):
    data = load_business_data(DATA_FILE)
    kpis = calculate_all_kpis(data)

    output_file = tmp_path / "kpi_summary.csv"

    result = save_kpi_summary(
        kpis,
        output_file,
    )

    assert Path(result).exists()

    content = output_file.read_text()

    assert "KPI,Value" in content
    assert "Total Revenue" in content
    assert "Net Revenue" in content
